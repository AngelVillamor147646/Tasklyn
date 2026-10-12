"""
Tasklyn — Lesson Question Generator (offline, rule-based)
==========================================================
Reads a lesson file (PDF, DOCX, TXT) and turns definitions and
glossary-style lines into identification and multiple-choice questions.
No internet connection or AI service is used.
"""
from __future__ import annotations

import os
import random
import re
from typing import Dict, List, Optional, Tuple

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".txt", ".md")

_STOP_STARTS = {
    "this", "that", "these", "those", "it", "its", "there", "here", "which",
    "who", "whom", "he", "she", "they", "we", "you", "i", "what", "when",
    "where", "why", "how", "also", "note", "example", "examples", "tip",
    "remember", "summary", "objective", "objectives", "figure", "table",
    "chapter", "lesson", "answer", "question", "however", "therefore",
    "another", "one", "some", "many", "most", "all", "each", "both",
}
_ARTICLES = ("the ", "a ", "an ")

_BULLET = re.compile(
    r"^\s*(?:[\u2022\u25cf\u25aa\u00b7*\-\u2013\u2014]|\(?\d{1,2}[.)]|\(?[a-zA-Z][.)])\s+"
)

# "Term: definition"  or  "Term - definition"  (one line, glossary style)
_GLOSSARY = re.compile(
    r"^(?P<term>[A-Za-z][\w\u2019' /]{1,50}?)\s*(?::|\s[-\u2013\u2014]\s)\s*(?P<rest>.{20,})$"
)

# Definition sentences: "X is defined as ...", "X refers to ...", "X is a ..."
_DEF_PATTERNS = [
    re.compile(
        r"^(?P<term>[^,;:()]{2,60}?)\s+(?:is|are)\s+defined\s+as\s+(?P<rest>.+)$", re.I),
    re.compile(
        r"^(?P<term>[^,;:()]{2,60}?)\s+(?:refers?\s+to|means?)\s+(?P<rest>.+)$", re.I),
    re.compile(
        r"^(?P<term>[^,;:()]{2,60}?)\s+(?:is|are)\s+"
        r"(?P<rest>(?:a|an|the|any|one|those|when|where|that|used|made|composed|located)\b.+)$",
        re.I),
]


# --------------------------------------------------------------------------
# Reading files
# --------------------------------------------------------------------------

def extract_text(path: str) -> str:
    """Return the plain text of a .pdf, .docx, .txt or .md file."""
    ext = os.path.splitext(path)[1].lower()

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    if ext == ".docx":
        import docx
        document = docx.Document(path)
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    parts.append(cell.text)
        return "\n".join(parts)

    if ext in (".txt", ".md"):
        with open(path, "r", encoding="utf-8-sig", errors="ignore") as f:
            return f.read()

    raise ValueError(
        f"Unsupported file type '{ext or 'unknown'}'. Use PDF, DOCX or TXT."
    )


# --------------------------------------------------------------------------
# Cleaning and splitting
# --------------------------------------------------------------------------

def _clean_lines(text: str) -> List[str]:
    lines: List[str] = []
    for raw in text.replace("\r", "\n").split("\n"):
        line = _BULLET.sub("", raw).strip()
        line = re.sub(r"\s+", " ", line)
        if not line:
            continue
        if re.fullmatch(r"(?:page\s*)?\d{1,4}", line, re.I):  # page numbers
            continue
        lines.append(line)
    return lines


def _sentences(lines: List[str]) -> List[str]:
    buf: List[str] = []
    for line in lines:
        # A short line with no punctuation is a heading: end it with a period
        if len(line.split()) <= 8 and not re.search(r"[.!?:;,]$", line):
            buf.append(line + ".")
        else:
            buf.append(line)
    text = " ".join(buf)
    parts = re.split(r"(?<=[.!?])\s+(?=[A-Z0-9\"\u201c(])", text)
    return [p.strip() for p in parts if p.strip()]


def _clean_term(term: str) -> str:
    t = term.strip(" \t\"'\u201c\u201d\u2018\u2019*_")
    low = t.lower()
    for article in _ARTICLES:
        if low.startswith(article):
            t = t[len(article):]
            break
    return t.strip()


def _term_ok(term: str) -> bool:
    if not term:
        return False
    words = term.split()
    if not 1 <= len(words) <= 5:
        return False
    if words[0].lower() in _STOP_STARTS:
        return False
    if not 3 <= len(term) <= 50:
        return False
    if re.search(r"\d{3,}", term):
        return False
    return bool(re.search(r"[A-Za-z]", term))


def _blank_out(text: str, term: str) -> str:
    """Replace every occurrence of the term with a blank."""
    return re.sub(r"\b" + re.escape(term) + r"\b", "_____", text, flags=re.I)


# --------------------------------------------------------------------------
# Finding (term, description) pairs
# --------------------------------------------------------------------------

def _find_pairs(text: str) -> List[Tuple[str, str]]:
    lines = _clean_lines(text)
    pairs: List[Tuple[str, str]] = []
    seen = set()

    def add(term: str, description: str) -> None:
        key = term.lower()
        if key in seen:
            return
        seen.add(key)
        pairs.append((term, description))

    # 1) Glossary-style lines:  "Mitosis: the process of cell division"
    for line in lines:
        m = _GLOSSARY.match(line)
        if not m:
            continue
        term = _clean_term(m.group("term"))
        rest = m.group("rest").strip()
        if _term_ok(term) and len(rest.split()) >= 4:
            add(term, _blank_out(rest, term))

    # 2) Definition sentences:  "Photosynthesis is the process by which..."
    for sentence in _sentences(lines):
        if not 40 <= len(sentence) <= 300:
            continue
        for pattern in _DEF_PATTERNS:
            m = pattern.match(sentence)
            if not m:
                continue
            term = _clean_term(m.group("term"))
            if not _term_ok(term) or len(m.group("rest").split()) < 4:
                continue
            description = "_____" + sentence[m.end("term"):]
            add(term, _blank_out(description, term))
            break

    return pairs


# --------------------------------------------------------------------------
# Public API
# --------------------------------------------------------------------------

def generate_questions(
    text: str,
    max_questions: int = 30,
    seed: Optional[int] = None,
) -> List[Dict]:
    """
    Turn lesson text into question dicts:
        {"question": str, "answer": str,
         "q_type": "identification" | "multiple_choice",
         "choices": [str, ...]}        # empty list for identification
    """
    rng = random.Random(seed)
    pairs = _find_pairs(text)
    all_terms = [term for term, _ in pairs]
    questions: List[Dict] = []

    for i, (term, description) in enumerate(pairs[:max_questions]):
        # Every second question is multiple choice (needs 3 other terms)
        if i % 2 == 1 and len(all_terms) >= 4:
            others = [t for t in all_terms if t.lower() != term.lower()]
            others.sort(key=lambda t: abs(len(t) - len(term)) + rng.random() * 3)
            distractors = rng.sample(others[:6], 3)
            choices = distractors + [term]
            rng.shuffle(choices)
            questions.append({
                "question": "Which term fits this description? " + description,
                "answer": term,
                "q_type": "multiple_choice",
                "choices": choices,
            })
        else:
            questions.append({
                "question": "Identify the term: " + description,
                "answer": term,
                "q_type": "identification",
                "choices": [],
            })
    return questions


def generate_from_file(
    path: str, max_questions: int = 30, seed: Optional[int] = None
) -> List[Dict]:
    """Read a lesson file and generate questions. Raises ValueError with a
    friendly message if the file can't be used."""
    text = extract_text(path)
    if not text.strip():
        raise ValueError(
            "No readable text found. If this is a scanned PDF (photos of "
            "pages), it can't be read."
        )
    questions = generate_questions(text, max_questions, seed)
    if not questions:
        raise ValueError(
            "Couldn't find any definitions. Lessons with sentences like "
            "'X is the ...' or lines like 'Term: meaning' work best."
        )
    return questions